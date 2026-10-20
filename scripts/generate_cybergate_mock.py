#!/usr/bin/env python3
"""Generate fictional CyberGate manager-dashboard data: 300 employees across 8 tables.

Output : data/cybergate-mock-data.xlsx  (8 sheets, original column names, in original order)
Run    : python3 scripts/generate_cybergate_mock.py
Needs  : openpyxl  (pip install openpyxl)

Every person, email and company here is made up. The random seed is fixed, so running the
script again gives exactly the same workbook. This file does not touch data/mock-data.xlsx.
"""
import datetime as dt
import math
import random
import unicodedata
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SEED = 2026
rng = random.Random(SEED)
TODAY = dt.date(2026, 10, 20)
DOMAIN = "acme-corp.example"
OUT = Path(__file__).resolve().parent.parent / "data" / "cybergate-mock-data.xlsx"


def rnd(x):
    return int(math.floor(x + 0.5))


def clamp(x, lo, hi):
    return max(lo, min(hi, x))


def pick_weighted(pairs):
    total = sum(w for _, w in pairs)
    r = rng.uniform(0, total)
    acc = 0
    for item, w in pairs:
        acc += w
        if r <= acc:
            return item
    return pairs[-1][0]


def ascii_slug(s):
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode()
    return "".join(c for c in s.lower() if c.isalnum() or c == "-")


def fmt_hours(h):
    whole = int(h)
    mins = rnd((h - whole) * 60)
    if mins == 60:
        whole, mins = whole + 1, 0
    return f"{whole}h" if mins == 0 else f"{whole}h {mins}m"


# --------------------------------------------------------------------------------------
# Competency framework (Tables 3 and 4)
# --------------------------------------------------------------------------------------
LEAD_FW = "CyberGate Leadership Framework"
TECH_FW = "CyberGate Technical Framework"

CATEGORIES = {
    "COM": (LEAD_FW, "Communication & Influence"),
    "LEA": (LEAD_FW, "Leadership & People"),
    "THK": (LEAD_FW, "Thinking & Decisions"),
    "LRN": (LEAD_FW, "Learning & Adaptability"),
    "DEL": (LEAD_FW, "Delivery & Ownership"),
    "DEF": (TECH_FW, "Defensive Operations"),
    "OFF": (TECH_FW, "Offensive & Application Security"),
    "INF": (TECH_FW, "Infrastructure Security"),
    "GRC": (TECH_FW, "Governance, Risk & Compliance"),
    "ENG": (TECH_FW, "Security Engineering"),
}

# code: (name, description, supporting text, [behavior phrase 1, behavior phrase 2], [module 1, module 2])
COMPETENCIES = {
    "COM-01": ("Communication",
               "Shares information clearly and concisely, in writing and in person, so others understand and can act.",
               "Assessed through written incident updates, reports and day-to-day team communication.",
               ["explain technical issues in plain language", "write clear, concise updates and reports"],
               ["Clear writing for security teams", "Explaining technical risk in plain language"]),
    "COM-02": ("Presentation Delivery",
               "Delivers well-structured presentations and briefings with confidence and composure.",
               "Assessed through briefings, demos and the presentation exercises in leadership programs.",
               ["present findings in a clear, logical structure", "handle questions from the audience with composure"],
               ["Presenting incident findings with confidence", "Structuring a 10-minute briefing"]),
    "COM-03": ("Stakeholder Management",
               "Builds trust with people outside the team and keeps them informed and aligned.",
               "Assessed through incident briefings, project updates and cross-team meetings.",
               ["identify who is affected by an issue and keep them informed", "adapt the message to each stakeholder's needs"],
               ["Stakeholder mapping and communication plans", "Managing expectations during incidents"]),
    "COM-04": ("Strategic Influence",
               "Gains support for ideas and decisions across the organisation, without relying on formal authority.",
               "Critical for lead and manager roles; assessed through proposals and steering discussions.",
               ["build support for security proposals", "influence decisions without formal authority"],
               ["Influencing without authority", "Building the business case for security"]),
    "COM-05": ("Relationship Building",
               "Listens actively and builds respectful, lasting working relationships.",
               "Assessed through peer feedback and observed group work.",
               ["build trust through active listening", "maintain respectful working relationships"],
               ["Active listening and trust building", "Working across teams and vendors"]),
    "LEA-01": ("Leadership",
               "Sets direction, acts as a role model and takes charge when situations escalate.",
               "Required from lead roles upward; assessed in interviews and leadership programs.",
               ["set direction and act as a visible role model", "take charge calmly when a situation escalates"],
               ["Stepping into a lead role", "Leading through a major incident"]),
    "LEA-02": ("Team Management",
               "Organises the team's work, develops people and handles conflict constructively.",
               "Required from lead roles upward; assessed through team feedback and interviews.",
               ["assign work fairly and coach team members", "resolve conflict within the team"],
               ["Coaching and developing junior analysts", "Running effective team rituals"]),
    "LEA-03": ("Performance Management",
               "Sets clear goals, follows through on commitments and gives useful feedback.",
               "Assessed through goal records, one-to-ones and follow-through on actions.",
               ["set clear goals and follow through on commitments", "give timely, specific feedback on performance"],
               ["Setting goals and giving feedback", "Holding others to account fairly"]),
    "LEA-04": ("Cross-Team Integration",
               "Works effectively with other teams and shares information and resources across boundaries.",
               "Assessed through joint projects, exercises and cross-functional rotations.",
               ["work effectively with teams outside their own", "share information and resources across teams"],
               ["Collaborating across security and IT", "Running joint exercises with other teams"]),
    "THK-01": ("Analytical Thinking",
               "Breaks complex problems into parts, reasons logically and reflects on outcomes.",
               "Assessed through investigations, case exercises and post-incident reviews.",
               ["break complex problems into logical parts", "reflect on outcomes and draw lessons"],
               ["Structured analysis for investigations", "Reflective practice and lessons learned"]),
    "THK-02": ("Problem Solving",
               "Defines the real problem, finds root causes and works out practical solutions.",
               "Assessed through troubleshooting, root cause analyses and case exercises.",
               ["define the real problem before acting", "find root causes rather than symptoms"],
               ["Root cause analysis techniques", "Framing problems before solving them"]),
    "THK-03": ("Strategic Thinking",
               "Connects daily work to long-term goals and generates new ideas.",
               "Required mainly for senior and manager roles; assessed through roadmaps and proposals.",
               ["connect daily work to long-term security goals", "generate and test new ideas"],
               ["Security strategy and roadmapping", "Thinking two steps ahead: scenario planning"]),
    "THK-04": ("Decision Making",
               "Makes sound, timely decisions by weighing risks and trade-offs.",
               "Assessed through incident decisions, prioritisation and scenario exercises.",
               ["make sound decisions with incomplete information", "weigh risks and trade-offs before deciding"],
               ["Decisions under uncertainty", "Risk-based prioritisation"]),
    "LRN-01": ("Continuous Learning",
               "Seeks feedback and new knowledge, learns quickly and applies what is learned.",
               "Assessed through learning activity, certification progress and response to feedback.",
               ["seek out feedback and new knowledge", "apply feedback quickly to improve"],
               ["Building a personal learning routine", "Learning from feedback"]),
    "LRN-02": ("Tolerance for Uncertainty",
               "Stays calm and effective under pressure and adapts when circumstances change.",
               "Assessed through behaviour during incidents, exercises and change.",
               ["stay calm and effective under pressure", "adapt plans when circumstances change"],
               ["Staying composed under pressure", "Working with incomplete information"]),
    "DEL-01": ("Operational Effectiveness",
               "Takes initiative and delivers work on time and to standard.",
               "Assessed through delivery records, ticket data and manager feedback.",
               ["take initiative to improve how work gets done", "deliver work on time and to standard"],
               ["Prioritising work and managing time", "Taking initiative on process improvements"]),
    "DEF-01": ("Threat Detection & Monitoring",
               "Builds, tunes and runs detection to spot malicious activity early.",
               "Assessed through detection rules authored, alert quality and lab exercises.",
               ["build and tune detection rules", "triage alerts and separate true from false positives"],
               ["SIEM use cases and alert tuning", "Detection engineering with Sigma rules"]),
    "DEF-02": ("Incident Response",
               "Contains, eradicates and recovers from security incidents using sound playbooks.",
               "Assessed through incident records, tabletop exercises and simulations.",
               ["contain and remediate security incidents", "follow and improve incident playbooks"],
               ["Incident handling playbooks", "Tabletop exercise practice"]),
    "DEF-03": ("Digital Forensics",
               "Collects and analyses digital evidence to reconstruct what happened.",
               "Assessed through forensic case work and practical labs.",
               ["collect and preserve digital evidence", "reconstruct attacker activity from artefacts"],
               ["Forensic triage on Windows endpoints", "Evidence handling and chain of custody"]),
    "DEF-04": ("Threat Intelligence",
               "Turns threat data into intelligence the organisation can act on.",
               "Assessed through intelligence reports and their use in detection and decisions.",
               ["turn threat data into actionable intelligence", "map adversary behaviour to MITRE ATT&CK"],
               ["Threat intelligence with MITRE ATT&CK", "Writing actionable intel reports"]),
    "OFF-01": ("Penetration Testing",
               "Plans and runs authorised tests that find exploitable weaknesses.",
               "Assessed through test reports, lab results and peer review.",
               ["plan and execute penetration tests", "report findings with clear remediation advice"],
               ["Web application testing labs", "Internal network penetration testing"]),
    "OFF-02": ("Secure Development",
               "Builds security into software design, code and delivery pipelines.",
               "Assessed through code reviews, threat models and pipeline controls.",
               ["review code for security flaws", "build security into the development lifecycle"],
               ["Secure coding for OWASP Top 10", "Threat modelling for developers"]),
    "OFF-03": ("Vulnerability Management",
               "Finds, prioritises and tracks vulnerabilities through to remediation.",
               "Assessed through scan coverage, remediation times and prioritisation quality.",
               ["prioritise vulnerabilities by real risk", "track remediation through to closure"],
               ["Vulnerability prioritisation with CVSS and EPSS", "Patch and remediation tracking"]),
    "INF-01": ("Network Security",
               "Designs and operates secure networks, segmentation and access rules.",
               "Assessed through design reviews, rule audits and lab work.",
               ["design and review secure network segmentation", "manage firewall and network access rules"],
               ["Network segmentation design", "Firewall rule review"]),
    "INF-02": ("Cloud Security",
               "Secures cloud workloads, configurations and identities.",
               "Assessed through posture findings, architecture reviews and labs.",
               ["secure cloud workloads and configurations", "detect and fix cloud misconfigurations"],
               ["Cloud security posture management", "Securing Kubernetes workloads"]),
    "INF-03": ("Identity & Access Management",
               "Manages identities, privileged access and authentication controls.",
               "Assessed through access reviews, PAM coverage and design work.",
               ["manage identities and privileged access", "design strong authentication and access controls"],
               ["Privileged access management", "Conditional access and MFA design"]),
    "INF-04": ("Endpoint Security",
               "Protects and hardens devices against common attacks.",
               "Assessed through EDR coverage, hardening baselines and incident data.",
               ["deploy and tune endpoint protection", "harden devices against common attacks"],
               ["EDR deployment and tuning", "Endpoint hardening baselines"]),
    "GRC-01": ("Risk Assessment",
               "Identifies, assesses and communicates security risk in business terms.",
               "Assessed through risk assessments delivered and the quality of risk reporting.",
               ["identify, assess and rate security risks", "communicate risk in business terms"],
               ["Risk assessment workshops", "Quantifying cyber risk"]),
    "GRC-02": ("Security Compliance",
               "Maps controls to laws, regulations and standards and evidences them.",
               "Assessed through control mappings, audit outcomes and evidence quality.",
               ["map controls to regulatory requirements", "prepare evidence for compliance reviews"],
               ["Mapping controls to ISO 27001 and NIS2", "Preparing for regulatory audits"]),
    "GRC-03": ("Security Policy & Audit",
               "Writes workable security policies and plans and runs audits.",
               "Assessed through policies authored and audits completed.",
               ["write and maintain security policies", "plan and run security audits"],
               ["Writing security policies people follow", "Audit fieldwork and evidence"]),
    "ENG-01": ("Scripting & Automation",
               "Automates security tasks and builds tooling with scripts and playbooks.",
               "Assessed through scripts, SOAR playbooks and time saved.",
               ["automate repetitive security tasks with scripts", "build and maintain security tooling"],
               ["Python for security automation", "SOAR playbook development"]),
    "ENG-02": ("Security Architecture",
               "Designs secure systems and reviews designs against security principles.",
               "Assessed through architecture decisions and design reviews.",
               ["design secure system architectures", "review designs against security principles"],
               ["Security architecture patterns", "Zero trust reference architecture"]),
}
LEAD_CODES = [c for c in COMPETENCIES if CATEGORIES[c[:3]][0] == LEAD_FW]
TECH_CODES = [c for c in COMPETENCIES if CATEGORIES[c[:3]][0] == TECH_FW]


def scale(code):
    """Leadership competencies use 1-5; technical competencies use 0-5 (0 = no exposure)."""
    return (1, 5) if code in LEAD_CODES else (0, 5)


LEVEL_TEMPLATES = {
    1: "With guidance, begins to {p}",
    2: "Is able to {p} in routine situations",
    3: "Is able to {p} independently in complex situations",
    4: "Helps others {p} and sets the standard for the team",
    5: "Shapes how the organisation is expected to {p}",
}

# --------------------------------------------------------------------------------------
# Organisation
# --------------------------------------------------------------------------------------
# Each section: (division, department, section, job role family, staff count, lead title,
#                [(role, grade, weight)], [technical competency codes, first two are core])
SECTIONS = [
    ("Information Security", "Security Operations", "Blue Team", "Security Operations", 34, "SOC Team Lead",
     [("SOC Analyst I", 1, 4), ("SOC Analyst II", 2, 4), ("Senior SOC Analyst", 3, 2)],
     ["DEF-01", "DEF-02", "DEF-04", "INF-01", "INF-04", "ENG-01"]),
    ("Information Security", "Security Operations", "Threat Intelligence", "Security Operations", 12, "Threat Intelligence Lead",
     [("Junior Threat Analyst", 1, 1), ("Threat Intelligence Analyst", 2, 3), ("Senior Threat Intelligence Analyst", 3, 2)],
     ["DEF-04", "DEF-01", "DEF-03", "GRC-01", "ENG-01"]),
    ("Information Security", "Security Operations", "Incident Response", "Security Operations", 20, "Incident Response Lead",
     [("Incident Responder", 2, 3), ("Digital Forensics Analyst", 2, 2), ("Senior Incident Responder", 3, 2)],
     ["DEF-02", "DEF-03", "DEF-01", "INF-04", "ENG-01"]),
    ("Information Security", "Offensive Security", "Red Team", "Security Engineering", 12, "Red Team Lead",
     [("Penetration Tester", 2, 3), ("Senior Penetration Tester", 3, 2), ("Red Team Operator", 3, 1)],
     ["OFF-01", "OFF-03", "INF-01", "ENG-01", "OFF-02"]),
    ("Information Security", "Offensive Security", "Application Security", "Security Engineering", 17, "Application Security Lead",
     [("Security Engineer", 2, 2), ("Application Security Engineer", 2, 3), ("Senior Application Security Engineer", 3, 2)],
     ["OFF-02", "OFF-03", "OFF-01", "ENG-01", "ENG-02", "INF-02"]),
    ("Information Security", "Identity & Access Management", "IAM Engineering", "Identity & Access", 14, "IAM Engineering Lead",
     [("IAM Analyst", 1, 2), ("IAM Engineer", 2, 3), ("Senior IAM Engineer", 3, 2)],
     ["INF-03", "INF-02", "ENG-01", "ENG-02", "GRC-02"]),
    ("Information Security", "Identity & Access Management", "Access Governance", "Identity & Access", 10, "Access Governance Lead",
     [("Access Governance Analyst", 1, 3), ("Senior Access Governance Analyst", 3, 2)],
     ["INF-03", "GRC-02", "GRC-03", "GRC-01"]),
    ("Technology & Infrastructure", "Cloud Platform", "Cloud Engineering", "Infrastructure", 28, "Cloud Engineering Lead",
     [("Junior Cloud Engineer", 1, 1), ("Cloud Engineer", 2, 4), ("Cloud Security Engineer", 2, 2), ("Senior Cloud Engineer", 3, 2)],
     ["INF-02", "ENG-01", "ENG-02", "INF-01", "INF-03", "OFF-03"]),
    ("Technology & Infrastructure", "Cloud Platform", "Platform Reliability", "Infrastructure", 14, "Platform Reliability Lead",
     [("Systems Administrator", 1, 2), ("Site Reliability Engineer", 2, 3), ("Senior Site Reliability Engineer", 3, 2)],
     ["INF-02", "ENG-01", "INF-04", "OFF-03", "DEF-01"]),
    ("Technology & Infrastructure", "Network Engineering", "Network Operations", "Infrastructure", 19, "Network Operations Lead",
     [("Network Technician", 1, 2), ("Network Engineer", 2, 3), ("Senior Network Engineer", 3, 2)],
     ["INF-01", "ENG-01", "DEF-01", "OFF-03"]),
    ("Technology & Infrastructure", "Network Engineering", "Network Security", "Infrastructure", 12, "Network Security Lead",
     [("Firewall Administrator", 1, 1), ("Network Security Engineer", 2, 3), ("Senior Network Security Engineer", 3, 2)],
     ["INF-01", "DEF-01", "OFF-03", "ENG-02", "INF-02"]),
    ("Technology & Infrastructure", "End-User Services", "Service Desk", "Infrastructure", 28, "Service Desk Lead",
     [("Service Desk Analyst", 1, 5), ("Senior Service Desk Analyst", 2, 2), ("Service Desk Coordinator", 2, 1)],
     ["INF-04", "INF-03", "DEF-02", "GRC-02"]),
    ("Technology & Infrastructure", "End-User Services", "Endpoint Engineering", "Infrastructure", 12, "Endpoint Engineering Lead",
     [("Desktop Support Technician", 1, 2), ("Endpoint Engineer", 2, 3), ("Endpoint Security Specialist", 2, 2)],
     ["INF-04", "OFF-03", "ENG-01", "DEF-01"]),
    ("Governance, Risk & Compliance", "Risk Management", "Enterprise Risk", "Governance, Risk & Compliance", 10, "Enterprise Risk Lead",
     [("GRC Analyst", 1, 2), ("Risk Analyst", 2, 3), ("Senior Risk Analyst", 3, 2)],
     ["GRC-01", "GRC-02", "GRC-03", "ENG-02"]),
    ("Governance, Risk & Compliance", "Risk Management", "Third-Party Risk", "Governance, Risk & Compliance", 8, "Third-Party Risk Lead",
     [("Vendor Risk Associate", 1, 2), ("Third-Party Risk Analyst", 2, 3)],
     ["GRC-01", "GRC-02", "GRC-03", "INF-02"]),
    ("Governance, Risk & Compliance", "Compliance & Audit", "Regulatory Compliance", "Governance, Risk & Compliance", 13, "Compliance Lead",
     [("Compliance Associate", 1, 2), ("Compliance Officer", 2, 3), ("Senior Compliance Officer", 3, 2)],
     ["GRC-02", "GRC-03", "GRC-01", "INF-03"]),
    ("Governance, Risk & Compliance", "Compliance & Audit", "Internal Audit", "Governance, Risk & Compliance", 9, "IT Audit Lead",
     [("IT Auditor", 2, 3), ("Senior IT Auditor", 3, 2)],
     ["GRC-03", "GRC-02", "GRC-01", "INF-03"]),
]
DEPARTMENT_HEADS = {
    "Security Operations": "Head of Security Operations",
    "Offensive Security": "Head of Offensive Security",
    "Identity & Access Management": "Head of Identity & Access Management",
    "Cloud Platform": "Head of Cloud Platform",
    "Network Engineering": "Head of Network Engineering",
    "End-User Services": "Head of End-User Services",
    "Risk Management": "Head of Risk Management",
    "Compliance & Audit": "Head of Compliance & Audit",
}
DIVISION_HEADS = {
    "Information Security": ("Chief Information Security Officer", ["DEF-02", "ENG-02", "GRC-01", "INF-03"]),
    "Technology & Infrastructure": ("Director of Technology & Infrastructure", ["INF-02", "ENG-02", "INF-01", "GRC-01"]),
    "Governance, Risk & Compliance": ("Director of Governance, Risk & Compliance", ["GRC-01", "GRC-02", "GRC-03", "ENG-02"]),
}
CEO = ("Victoria Ashdown", f"victoria.ashdown@{DOMAIN}")  # outside the 300, only appears as a manager

DEPARTMENT_NEEDS = {
    "Security Operations": ("AI-assisted alert triage and SOAR playbooks", 12),
    "Offensive Security": ("Testing AI and LLM applications for prompt injection", 10),
    "Identity & Access Management": ("Passwordless authentication rollout", 8),
    "Cloud Platform": ("Cloud incident response and forensics", 12),
    "Network Engineering": ("Zero trust network access migration", 10),
    "End-User Services": ("Recognising and escalating social engineering", 6),
    "Risk Management": ("NIS2 and DORA third-party risk requirements", 8),
    "Compliance & Audit": ("Auditing AI systems under the EU AI Act", 10),
}

FIRST_NAMES = [
    "Amina", "Emir", "Lejla", "Adnan", "Selma", "Haris", "Ajla", "Tarik", "Dina", "Kenan", "Merima", "Amar",
    "Sara", "Daniel", "Omar", "Layla", "Yusuf", "Fatima", "Karim", "Noor", "Hassan", "Mariam", "Zayed", "Huda",
    "Priya", "Arjun", "Ananya", "Rohan", "Kavya", "Vikram", "Meera", "Sanjay", "Neha", "Aditya",
    "Maria", "Jose", "Ana", "Miguel", "Isabel", "Carlos", "Lucia", "Diego", "Elena", "Pablo",
    "Emma", "Liam", "Olivia", "Noah", "Sophie", "James", "Chloe", "Ethan", "Grace", "Lucas", "Hannah", "Ben",
    "Wei", "Mei", "Jun", "Yuki", "Hiro", "Min-jun", "Ji-woo", "Lin", "Kenji", "Aiko",
    "Chinedu", "Amara", "Kofi", "Ama", "Tunde", "Zainab", "Kwesi", "Nia", "Femi", "Imani",
    "Lukas", "Anna", "Jonas", "Lea", "Felix", "Marta", "Piotr", "Kasia", "Ivan", "Nina", "Marko", "Ivana",
    "Mateo", "Valentina", "Rafael", "Camila", "Andrei", "Ioana", "Nikos", "Eleni", "Mehmet", "Ayse",
    "Joseph", "Rhea", "Paolo", "Giulia", "Sven", "Freya", "Aisha", "Rami", "Dalia", "Samir",
]
LAST_NAMES = [
    "Hodzic", "Begic", "Kovac", "Delic", "Mehic", "Hadzic", "Ibrahimovic", "Softic", "Alic", "Mujic", "Pasic",
    "Haddad", "Mansour", "Al-Farsi", "Khoury", "Nasser", "Saleh", "Rahman", "Aziz", "Hamdan", "Qureshi",
    "Sharma", "Patel", "Iyer", "Nair", "Reddy", "Gupta", "Menon", "Rao", "Kapoor", "Joshi",
    "Santos", "Reyes", "Cruz", "Garcia", "Fernandez", "Lopez", "Morales", "Ortega", "Navarro", "Romero",
    "Smith", "Reed", "Walker", "Bennett", "Clarke", "Hughes", "Turner", "Price", "Ward", "Foster", "Murphy",
    "Chen", "Wang", "Tanaka", "Sato", "Kim", "Park", "Nguyen", "Tran", "Liu", "Watanabe",
    "Okafor", "Mensah", "Adeyemi", "Boateng", "Okonkwo", "Diallo", "Owusu", "Eze", "Abubakar", "Kamara",
    "Muller", "Schmidt", "Weber", "Novak", "Horvat", "Kowalski", "Nowak", "Petrov", "Popescu", "Jovanovic",
    "Rossi", "Bianchi", "Papadopoulos", "Yilmaz", "Demir", "Andersen", "Larsen", "Lindqvist", "Dubois", "Moreau",
    "Silva", "Costa", "Pereira", "Oliveira", "Martins", "Brennan", "Gallagher", "O'Neill", "Fitzgerald", "Doyle",
]

used_names, used_emails = set(), set()


def new_person_name():
    while True:
        name = f"{rng.choice(FIRST_NAMES)} {rng.choice(LAST_NAMES)}"
        if name not in used_names:
            used_names.add(name)
            first, last = name.split(" ", 1)
            base = f"{ascii_slug(first)}.{ascii_slug(last)}"
            email, n = f"{base}@{DOMAIN}", 2
            while email in used_emails:
                email, n = f"{base}{n}@{DOMAIN}", n + 1
            used_emails.add(email)
            return name, email


# Archetypes give the data its spread: who is ready to move up, who is struggling, etc.
ARCHETYPES = {
    #               weight, soft offset, tech offset, growth/day in programs
    "rising_star":   (12, 0.85, 0.55, (0.25, 0.40)),
    "solid":         (38, 0.05, 0.10, (0.10, 0.25)),
    "tech_expert":   (15, -0.75, 0.90, (0.05, 0.20)),
    "developing":    (17, -0.45, -0.45, (0.10, 0.30)),
    "new_joiner":    (12, -0.70, -0.90, (0.15, 0.35)),
    "struggling":    (6, -1.15, -0.75, (-0.05, 0.10)),
}


def pick_archetype(grade):
    weights = []
    for name, (w, *_rest) in ARCHETYPES.items():
        if grade >= 4 and name in ("new_joiner", "struggling"):
            w = w * 0.25
        if grade >= 4 and name in ("rising_star", "solid"):
            w = w * 1.4
        weights.append((name, w))
    return pick_weighted(weights)


employees = []


def add_employee(role, grade, family, division, department, section, tech_codes, extended):
    name, email = new_person_name()
    arch = pick_archetype(grade)
    _, soft, tech, growth = ARCHETYPES[arch]
    emp = dict(
        name=name, email=email, role=role, grade=grade, family=family, division=division,
        department=department, section=section, tech_codes=tech_codes, extended=extended,
        archetype=arch, soft=soft + rng.gauss(0, 0.25), tech=tech + rng.gauss(0, 0.25),
        growth=rng.uniform(*growth), line_manager=None, team_name=None, team_manager=None,
    )
    employees.append(emp)
    return emp


def short_title(role):
    for suffix in (" I", " II", " III"):
        if role.endswith(suffix):
            return role[: -len(suffix)]
    return role


# Division heads, department heads, section leads, then staff
div_heads, dept_heads, section_leads = {}, {}, {}
for division, (title, codes) in DIVISION_HEADS.items():
    div_heads[division] = add_employee(title, 6, "Leadership", division, f"{division} Office", "Leadership Team",
                                       codes, f"{title} – {division}")
for sec in SECTIONS:
    division, department = sec[0], sec[1]
    if department not in dept_heads:
        title = DEPARTMENT_HEADS[department]
        codes = []
        for s in SECTIONS:
            if s[1] == department:
                codes += [c for c in s[7][:3] if c not in codes]
        dept_heads[department] = add_employee(title, 5, "Leadership", division, department, "Department Management",
                                              codes[:5], f"{title} – {division}")
for sec in SECTIONS:
    division, department, section, family, n_staff, lead_title, roles, codes = sec
    section_leads[section] = add_employee(lead_title, 4, family, division, department, section, codes,
                                          f"{lead_title} – {section}")
    specialisms = {
        "Blue Team": ["Threat Detection", "Monitoring", "Triage"], "Cloud Engineering": ["AWS", "Azure", "Kubernetes"],
        "Service Desk": ["Tier 1", "Tier 2", "VIP Support"], "Application Security": ["Web", "Mobile", "CI/CD"],
    }.get(section)
    for _ in range(n_staff):
        role, grade, _w = pick_weighted([((r, g, w), w) for r, g, w in roles])
        ext = f"{role} – {section}" + (f", {rng.choice(specialisms)}" if specialisms else "")
        add_employee(role, grade, family, division, department, section, codes, ext)

assert len(employees) == 300, len(employees)

# Reporting lines and teams
for e in employees:
    if e["grade"] == 6:
        e["line_manager"] = CEO
        e["team_name"], e["team_manager"] = "Executive Leadership", CEO
    elif e["grade"] == 5:
        h = div_heads[e["division"]]
        e["line_manager"] = (h["name"], h["email"])
        e["team_name"], e["team_manager"] = f"{e['division']} Leadership", (h["name"], h["email"])
    elif e["grade"] == 4:
        h = dept_heads[e["department"]]
        e["line_manager"] = (h["name"], h["email"])
        e["team_name"], e["team_manager"] = f"{e['department']} Leads", (h["name"], h["email"])
    else:
        lead = section_leads[e["section"]]
        e["team_name"], e["team_manager"] = e["section"], (lead["name"], lead["email"])
        # about 1 in 12 report on paper to the department head (matrix setup)
        if rng.random() < 0.08:
            h = dept_heads[e["department"]]
            e["line_manager"] = (h["name"], h["email"])
        else:
            e["line_manager"] = (lead["name"], lead["email"])

# --------------------------------------------------------------------------------------
# Required levels, priorities, actual levels
# --------------------------------------------------------------------------------------
PEOPLE_LEAD = {"LEA-01", "LEA-02", "LEA-03"}
STRATEGIC = {"COM-04", "THK-03"}
CORE_SOFT = {"COM-01", "THK-01", "THK-02", "LRN-01", "DEL-01"}


def required_level(code, grade, tech_codes):
    if code in LEAD_CODES:
        if code in PEOPLE_LEAD:
            table = {1: 1, 2: 2, 3: 2, 4: 4, 5: 4, 6: 5}
        elif code in STRATEGIC:
            table = {1: 1, 2: 2, 3: 3, 4: 3, 5: 4, 6: 5}
        elif code in CORE_SOFT:
            table = {1: 2, 2: 3, 3: 3, 4: 4, 5: 4, 6: 4}
        else:
            table = {1: 2, 2: 2, 3: 3, 4: 3, 5: 4, 6: 4}
        return table[grade]
    if code in tech_codes[:2]:
        return {1: 2, 2: 3, 3: 4, 4: 4, 5: 4, 6: 3}[grade]
    return {1: 1, 2: 2, 3: 3, 4: 3, 5: 3, 6: 3}[grade]


def priority(code, grade, tech_codes):
    if code in TECH_CODES:
        return "High" if code in tech_codes[:2] else "Medium"
    if grade >= 4:
        if code in PEOPLE_LEAD or code in ("THK-04", "COM-03") or (grade >= 5 and code in STRATEGIC):
            return "High"
        if code in ("COM-05", "LRN-01", "LRN-02"):
            return "Low"
        return "Medium"
    if code in ("COM-01", "THK-02", "DEL-01"):
        return "High"
    if code in ("COM-03", "THK-01", "LRN-01", "LEA-04", "COM-05"):
        return "Medium"
    return "Low"


for e in employees:
    e["codes"] = LEAD_CODES + e["tech_codes"]
    e["latent"] = {}
    for code in e["codes"]:
        req = required_level(code, e["grade"], e["tech_codes"])
        off = e["soft"] if code in LEAD_CODES else e["tech"]
        lo, hi = scale(code)
        e["latent"][code] = clamp(rnd(req + off + rng.gauss(0, 0.6)), lo, hi)

# --------------------------------------------------------------------------------------
# Table 1: assessment status
# --------------------------------------------------------------------------------------
for e in employees:
    if e["archetype"] == "new_joiner":
        sa = pick_weighted([("Completed", 55), ("In Progress", 25), ("Not Started", 20)])
    else:
        sa = pick_weighted([("Completed", 92), ("In Progress", 5), ("Not Started", 3)])
    if sa == "Completed":
        vi = pick_weighted([("Completed", 84), ("Scheduled", 11), ("Not Scheduled", 5)])
    else:
        vi = "Not Scheduled"
    if vi == "Completed":
        rep = pick_weighted([("Delivered", 88), ("In Draft", 12)])
    else:
        rep = "Pending"
    cft_w = [("Completed", 14), ("In Progress", 12), ("Recommended", 18), ("Not Required", 56)]
    if e["archetype"] == "rising_star":
        cft_w = [("Completed", 30), ("In Progress", 20), ("Recommended", 25), ("Not Required", 25)]
    cft = pick_weighted(cft_w) if rep == "Delivered" else pick_weighted([("Recommended", 10), ("Not Required", 90)])
    e.update(sa=sa, vi=vi, report=rep, cft=cft)

# --------------------------------------------------------------------------------------
# Training catalogue and sessions (Tables 5 and 6)
# --------------------------------------------------------------------------------------
# key: (title, type, mode, hours, days)
COURSES = {
    "awareness25": ("Security Awareness Annual Refresher 2025", "Compliance", "Self-paced", 1.5, 1),
    "awareness26": ("Security Awareness Annual Refresher 2026", "Compliance", "Self-paced", 1.5, 1),
    "gdpr": ("Data Protection & GDPR Essentials", "Compliance", "Online", 2, 1),
    "ai_sec": ("AI for Security Teams", "Technical", "Online", 6, 1),
    "ai_lit": ("Prompting & AI Safety at Work", "AI Literacy", "Online", 3, 1),
    "siem": ("SIEM Detection Engineering", "Technical", "Hybrid", 16, 2),
    "hunt": ("Threat Hunting Fundamentals", "Technical", "Online", 12, 2),
    "ir_ws": ("Incident Handling Workshop", "Technical", "In-person", 16, 2),
    "malware": ("Malware Analysis Basics", "Technical", "Online", 20, 3),
    "forensics": ("Digital Forensics & Evidence Handling", "Technical", "In-person", 24, 3),
    "webpt": ("Web Application Penetration Testing", "Technical", "Hybrid", 24, 3),
    "codereview": ("Secure Code Review", "Technical", "Online", 12, 2),
    "cloudattack": ("Cloud Attack Paths", "Technical", "Online", 16, 2),
    "threatmodel": ("Threat Modelling Workshop", "Technical", "In-person", 8, 1),
    "zt_id": ("Zero Trust Identity Architecture", "Technical", "Online", 12, 2),
    "pam": ("Privileged Access Management in Practice", "Technical", "Hybrid", 16, 2),
    "accessauto": ("Access Review Automation", "Technical", "Online", 8, 1),
    "cloudfound": ("Cloud Security Foundations", "Technical", "Online", 12, 2),
    "k8s": ("Kubernetes Security", "Technical", "Hybrid", 16, 2),
    "netseg": ("Network Segmentation & Firewall Policy", "Technical", "In-person", 16, 2),
    "iac": ("Infrastructure as Code Security", "Technical", "Online", 12, 2),
    "edr": ("Endpoint Hardening with EDR", "Technical", "Online", 10, 2),
    "itil": ("ITIL 4 Foundation", "Technical", "Online", 16, 2),
    "iso": ("ISO 27001 Implementation", "Technical", "Hybrid", 24, 3),
    "tpr": ("Third-Party Risk Assessment", "Technical", "Online", 8, 1),
    "nis2": ("NIS2 and DORA Readiness", "Compliance", "Online", 6, 1),
    "auditev": ("Audit Evidence & Sampling", "Technical", "In-person", 12, 2),
    "fair": ("Risk Quantification with FAIR", "Technical", "Online", 14, 2),
    "feedback": ("Giving and Receiving Feedback", "Soft Skills Workshop", "In-person", 4, 1),
    "time": ("Time & Priority Management", "Soft Skills Workshop", "Online", 3, 1),
    "execwriting": ("Writing for Executives", "Soft Skills Workshop", "Online", 4, 1),
    "coaching": ("Coaching Skills for Senior Staff", "Soft Skills Workshop", "In-person", 8, 1),
    "rot_cloud": ("Cloud Security Rotation", "Cross-Functional", "Hybrid", 40, 10),
    "rot_soc": ("SOC Shadowing Programme", "Cross-Functional", "In-person", 24, 5),
    "rot_grc": ("GRC for Engineers", "Cross-Functional", "Online", 12, 3),
    "rot_purple": ("Red Team / Blue Team Exchange", "Cross-Functional", "In-person", 16, 4),
    # observed programs (trainers score every day -> Table 6)
    "elb": ("Emerging Leaders Bootcamp", "Leadership Program", "In-person", 24, 3),
    "mgr": ("Manager Essentials", "Leadership Program", "Hybrid", 21, 3),
    "cwi": ("Communicate with Impact", "Soft Skills Workshop", "In-person", 14, 2),
    "sll": ("Strategic Leadership Lab", "Leadership Program", "In-person", 16, 2),
}
OBSERVED = {"elb", "mgr", "cwi", "sll"}
FAMILY_COURSES = {
    "Security Operations": ["siem", "hunt", "ir_ws", "malware", "forensics"],
    "Security Engineering": ["webpt", "codereview", "cloudattack", "threatmodel"],
    "Identity & Access": ["zt_id", "pam", "accessauto", "cloudfound"],
    "Infrastructure": ["cloudfound", "k8s", "netseg", "iac", "edr", "itil"],
    "Governance, Risk & Compliance": ["iso", "tpr", "nis2", "auditev", "fair"],
    "Leadership": ["fair", "nis2", "threatmodel", "cloudfound"],
}
CFT_COURSE = {
    "Security Operations": "rot_cloud", "Security Engineering": "rot_purple", "Identity & Access": "rot_grc",
    "Infrastructure": "rot_soc", "Governance, Risk & Compliance": "rot_grc", "Leadership": "rot_soc",
}
TRAINERS = {"Nadia Karim": 0.15, "Tomas Herrera": -0.20, "Ingrid Solberg": 0.0, "Kwame Asante": 0.30,
            "Elif Demir": -0.05, "Rahul Menon": 0.10}

PERIOD_START, PERIOD_END = dt.date(2025, 1, 6), dt.date(2026, 12, 14)


def monday_on_or_after(d):
    return d + dt.timedelta(days=(7 - d.weekday()) % 7)


def add_workdays(d, n):
    while n > 0:
        d += dt.timedelta(days=1)
        if d.weekday() < 5:
            n -= 1
    return d


def course_dates(key):
    """Start dates for a course's sessions. Observed programs get fixed, readable schedules."""
    fixed = {
        "elb": ["2025-03-10", "2025-06-09", "2025-10-13", "2026-02-09", "2026-05-11", "2026-09-14", "2026-10-19", "2026-11-16"],
        "mgr": ["2025-04-07", "2025-11-03", "2026-03-09", "2026-06-15", "2026-10-05", "2026-12-07"],
        "cwi": ["2025-02-17", "2025-05-19", "2025-09-22", "2026-01-26", "2026-04-20", "2026-07-06", "2026-10-26"],
        "sll": ["2025-09-08", "2026-04-13", "2026-11-23"],
        "awareness25": ["2025-01-13", "2025-03-03", "2025-05-05", "2025-09-01", "2025-11-03"],
        "awareness26": ["2026-01-12", "2026-03-02", "2026-05-04", "2026-09-07", "2026-11-02"],
    }
    if key in fixed:
        return [dt.date.fromisoformat(s) for s in fixed[key]]
    n = rng.randint(4, 9)
    span = (PERIOD_END - PERIOD_START).days
    return sorted({monday_on_or_after(PERIOD_START + dt.timedelta(days=rng.randint(0, span))) for _ in range(n)})


sessions = {}  # key -> list of session dicts
for key in COURSES:
    title, ttype, mode, hours, days = COURSES[key]
    sessions[key] = [dict(key=key, start=s, end=add_workdays(s, days - 1), trainer=rng.choice(list(TRAINERS)),
                          batch=f"Batch {i + 1:02d}", members=[]) for i, s in enumerate(course_dates(key))]


def cert_prep_session(cert):
    key = f"prep::{cert}"
    if key not in sessions:
        hours = 40 if cert in ("CISSP", "CISM", "CISA", "CRISC", "OSCP", "CCSP", "GIAC GCFA") else rng.choice([16, 20, 24])
        COURSES[key] = (f"{cert} Exam Prep", "Certification Prep", rng.choice(["Self-paced", "Online"]), hours, 30)
        sessions[key] = [dict(key=key, start=s, end=s + dt.timedelta(days=rng.randint(28, 56)),
                              trainer=None, batch=f"Batch {i + 1:02d}", members=[])
                         for i, s in enumerate(course_dates(key))]
    return key


# --------------------------------------------------------------------------------------
# Table 7: IDP certifications (built before training so cert prep can be scheduled)
# --------------------------------------------------------------------------------------
PROF_CERTS = {
    "Security Operations": ["CompTIA Security+", "CompTIA CySA+", "GIAC GCIH", "GIAC GCFA", "CISSP"],
    "Security Engineering": ["CompTIA Security+", "INE eJPT", "OSCP", "GIAC GWAPT", "CISSP"],
    "Identity & Access": ["CompTIA Security+", "Microsoft SC-300", "CISSP", "CCSP"],
    "Infrastructure": ["CompTIA Network+", "CompTIA Security+", "Microsoft AZ-500", "AWS Certified Security – Specialty", "CCSP"],
    "Governance, Risk & Compliance": ["ISO/IEC 27001 Lead Implementer", "ISO/IEC 27001 Lead Auditor", "CISA", "CRISC", "CISM"],
    "Leadership": ["CISM", "CISSP", "EC-Council CCISO"],
}
AI_CERTS = {
    "Governance, Risk & Compliance": ["IAPP AIGP", "Microsoft Azure AI Fundamentals (AI-900)"],
    "Leadership": ["Google Cloud Generative AI Leader", "IAPP AIGP"],
    "default": ["Microsoft Azure AI Fundamentals (AI-900)", "AWS Certified AI Practitioner",
                "Microsoft Azure AI Engineer Associate (AI-102)"],
}
CYCLES = [("2025 H2", dt.date(2025, 7, 1), dt.date(2025, 12, 31)),
          ("2026 H1", dt.date(2026, 1, 1), dt.date(2026, 6, 30)),
          ("2026 H2", dt.date(2026, 7, 1), dt.date(2026, 12, 31))]
DEFER_REASONS = ["project workload", "a change of role", "extended leave", "a higher-priority goal", "team staffing gaps"]
STUDY = ["the official course", "self-study and practice exams", "a vendor bootcamp", "the internal study group"]


def rand_date(a, b):
    return a + dt.timedelta(days=rng.randint(0, max(0, (b - a).days)))


def cert_status(current, arch):
    good = arch in ("rising_star", "tech_expert")
    if current:
        w = [("Not Started", 22), ("In Progress", 36), ("Exam Scheduled", 17), ("Achieved", 18), ("Deferred", 7)]
        if good:
            w = [("Not Started", 10), ("In Progress", 30), ("Exam Scheduled", 22), ("Achieved", 33), ("Deferred", 5)]
        if arch == "struggling":
            w = [("Not Started", 40), ("In Progress", 35), ("Exam Scheduled", 5), ("Achieved", 5), ("Deferred", 15)]
        return pick_weighted(w)
    w = [("Achieved", 60), ("Failed", 11), ("Deferred", 19), ("Not Started", 10)]
    if good:
        w = [("Achieved", 80), ("Failed", 6), ("Deferred", 10), ("Not Started", 4)]
    if arch == "struggling":
        w = [("Achieved", 30), ("Failed", 25), ("Deferred", 25), ("Not Started", 20)]
    return pick_weighted(w)


def cert_description(status, start, end, current):
    if status == "Achieved":
        when = rand_date(start, min(end, TODAY))
        return rng.choice([f"Passed on {when.isoformat()}.", f"Certified in {when.strftime('%B %Y')}; credential added to profile.",
                           f"Passed first time on {when.isoformat()}."]), when
    if status == "Exam Scheduled":
        when = rand_date(TODAY + dt.timedelta(days=7), end)
        return f"Exam booked for {when.isoformat()}.", when
    if status == "In Progress":
        return f"Studying via {rng.choice(STUDY)}; about {rng.choice([20, 35, 50, 60, 75, 85])}% of the syllabus covered.", None
    if status == "Failed":
        when = rand_date(start, end)
        return f"Did not pass on {when.isoformat()}; retake planned for next cycle.", when
    if status == "Deferred":
        return f"Deferred due to {rng.choice(DEFER_REASONS)}.", None
    return ("Not started; study plan to be agreed with line manager." if current
            else "Not started during this cycle."), None


idp_rows = []
for e in employees:
    fam = e["family"]
    ladder = PROF_CERTS[fam]
    if e["section"] == "Service Desk":
        ladder = ["CompTIA A+", "CompTIA Network+", "CompTIA Security+"]
    if e["section"] == "Application Security":
        ladder = ["CompTIA Security+", "GIAC GWAPT", "CSSLP", "CISSP"]
    ai_ladder = AI_CERTS.get(fam, AI_CERTS["default"])
    idx = clamp(e["grade"] - 1 + (1 if e["archetype"] == "tech_expert" else 0), 0, len(ladder) - 1)
    ai_idx = 0 if e["grade"] <= 3 else min(1, len(ai_ladder) - 1)
    cycles = CYCLES[2:] if e["archetype"] == "new_joiner" else (CYCLES[1:] if rng.random() < 0.25 else CYCLES)
    e["current_cert"] = None
    for name, start, end in cycles:
        current = name == "2026 H2"
        prof = ladder[min(idx, len(ladder) - 1)]
        st = cert_status(current, e["archetype"])
        desc, when = cert_description(st, start, end, current)
        has_ai = rng.random() > (0.12 if current else 0.3)
        ai = ai_ladder[min(ai_idx, len(ai_ladder) - 1)] if has_ai else None
        ai_st = cert_status(current, e["archetype"]) if has_ai else None
        ai_desc = cert_description(ai_st, start, end, current)[0] if has_ai else None
        idp_rows.append([name, e["email"], prof, st, desc, ai, ai_st, ai_desc])
        if current:
            e["current_cert"] = (prof, st)
        if st == "Achieved":
            idx += 1
        if has_ai and ai_st == "Achieved":
            ai_idx += 1

# --------------------------------------------------------------------------------------
# Table 5: enrolments
# --------------------------------------------------------------------------------------
training_rows = []
observed_enrolments = []  # (employee, session, status, attended_days)


def enrol(e, key, earliest=PERIOD_START):
    options = [s for s in sessions[key] if s["start"] >= earliest] or sessions[key]
    s = rng.choice(options)
    title, ttype, mode, hours, days = COURSES[key]
    if s["end"] < TODAY:
        no_show_bias = 2.5 if e["archetype"] == "struggling" else 1
        status = pick_weighted([("Completed", 88), ("Incomplete", 7 * no_show_bias), ("Cancelled", 4)])
    elif s["start"] <= TODAY:
        status = "In Progress"
    else:
        status = "Not Started"
    if status == "Completed":
        attendance = pick_weighted([("Attended", 91), ("Partial", 9)])
        credited = hours if attendance == "Attended" else round(hours * rng.uniform(0.5, 0.85) * 2) / 2
    elif status == "Incomplete":
        attendance = pick_weighted([("Partial", 70), ("No-Show", 30)])
        credited = round(hours * rng.uniform(0.15, 0.5) * 2) / 2 if attendance == "Partial" else 0
    elif status == "In Progress":
        attendance = "Attended"
        total = max(1, (s["end"] - s["start"]).days + 1)
        credited = round(hours * min(1, ((TODAY - s["start"]).days + 1) / total) * 2) / 2
    else:
        attendance, credited = "", 0
    s["members"].append(e["email"])
    training_rows.append(dict(session=s, email=e["email"], title=title, ttype=ttype, mode=mode,
                              attendance=attendance, status=status, hours=credited))
    if key in OBSERVED and status in ("Completed", "Incomplete", "In Progress"):
        observed_enrolments.append((e, s, status, attendance))


for e in employees:
    joiner = e["archetype"] == "new_joiner"
    start = dt.date(2026, 3, 2) if joiner else PERIOD_START
    if not joiner:
        enrol(e, "awareness25")
    enrol(e, "awareness26")
    if rng.random() < 0.35:
        enrol(e, "gdpr", start)
    tech_pool = FAMILY_COURSES[e["family"]][:]
    rng.shuffle(tech_pool)
    n_tech = rng.randint(0, 1) if joiner else rng.randint(1, 3)
    if e["archetype"] in ("rising_star", "tech_expert"):
        n_tech += 1
    for key in tech_pool[:n_tech]:
        enrol(e, key, start)
    if rng.random() < 0.45:
        enrol(e, rng.choice(["ai_sec", "ai_lit"]), start)
    if rng.random() < 0.4:
        soft = ["feedback", "time", "execwriting"] + (["coaching"] if e["grade"] >= 3 else [])
        enrol(e, rng.choice(soft), start)
    # observed programs
    p = {"rising_star": 0.85, "solid": 0.5, "tech_expert": 0.35, "developing": 0.45, "new_joiner": 0.25, "struggling": 0.45}[e["archetype"]]
    if rng.random() < p:
        if e["grade"] <= 1:
            enrol(e, "cwi", start)
        elif e["grade"] <= 3:
            enrol(e, rng.choice(["elb", "elb", "cwi"]), start)
        elif e["grade"] <= 5:
            enrol(e, "mgr", start)
        else:
            enrol(e, "sll", start)
    if e["grade"] in (3, 4) and rng.random() < 0.15:
        enrol(e, "elb" if e["grade"] == 3 else "sll", start)
    # cross-functional
    if e["cft"] in ("Completed", "In Progress"):
        key = CFT_COURSE[e["family"]]
        pool = [s for s in sessions[key] if (s["end"] < TODAY) == (e["cft"] == "Completed")
                and not (e["cft"] == "In Progress" and s["start"] > TODAY)]
        if not pool:  # make a session that fits the status
            s0 = TODAY - dt.timedelta(days=rng.randint(40, 300)) if e["cft"] == "Completed" else TODAY - dt.timedelta(days=rng.randint(1, 5))
            s0 = monday_on_or_after(s0 - dt.timedelta(days=6))
            days = COURSES[key][4]
            pool = [dict(key=key, start=s0, end=add_workdays(s0, days - 1), trainer=None,
                         batch=f"Batch {len(sessions[key]) + 1:02d}", members=[])]
            sessions[key].append(pool[0])
        s = rng.choice(pool)
        title, ttype, mode, hours, days = COURSES[key]
        st = e["cft"]
        credited = hours if st == "Completed" else round(hours * 0.4 * 2) / 2
        s["members"].append(e["email"])
        training_rows.append(dict(session=s, email=e["email"], title=title, ttype=ttype, mode=mode,
                                  attendance="Attended", status=st, hours=credited))
    # certification prep for the current cycle goal
    if e["current_cert"] and e["current_cert"][1] in ("In Progress", "Exam Scheduled", "Achieved"):
        enrol(e, cert_prep_session(e["current_cert"][0]), dt.date(2026, 5, 1))

# Training Record Codes: one per session, numbered by start date within the year
all_sessions = sorted({id(r["session"]): r["session"] for r in training_rows}.values(), key=lambda s: (s["start"], s["key"]))
year_seq = {}
for s in all_sessions:
    y = s["start"].year
    year_seq[y] = year_seq.get(y, 0) + 1
    s["code"] = f"TR-{y}-{year_seq[y]:04d}"
    if s["trainer"] is None:
        s["trainer"] = rng.choice(list(TRAINERS))

# --------------------------------------------------------------------------------------
# Table 6: observation scores
# --------------------------------------------------------------------------------------
BEHAVIORS = [
    ("Confidence / composure", "Presentation Delivery", "COM-02"),
    ("Accountability / follow-through", "Performance Management", "LEA-03"),
    ("Emotional regulation", "Tolerance for Uncertainty", "LRN-02"),
    ("Openness to feedback", "Continuous Learning", "LRN-01"),
    ("Clarity of message", "Communication", "COM-01"),
    ("Structure of message", "Presentation Delivery", "COM-02"),
    ("Active listening", "Relationship Building", "COM-05"),
    ("Adaptation to audience", "Stakeholder Management", "COM-03"),
    ("Conciseness", "Communication", "COM-01"),
    ("Participation", "Cross-Team Integration", "LEA-04"),
    ("Respectful interaction", "Relationship Building", "COM-05"),
    ("Influence in group", "Strategic Influence", "COM-04"),
    ("Conflict handling", "Team Management", "LEA-02"),
    ("Collaboration", "Cross-Team Integration", "LEA-04"),
    ("Problem definition", "Problem Solving", "THK-02"),
    ("Logical thinking", "Analytical Thinking", "THK-01"),
    ("Idea generation", "Strategic Thinking", "THK-03"),
    ("Decision quality", "Decision Making", "THK-04"),
    ("Initiative", "Operational Effectiveness", "DEL-01"),
    ("Ownership", "Performance Management", "LEA-03"),
    ("Guiding others", "Team Management", "LEA-02"),
    ("Presence / authority", "Leadership", "LEA-01"),
    ("Learning speed", "Continuous Learning", "LRN-01"),
    ("Applies feedback", "Continuous Learning", "LRN-01"),
    ("Reflection quality", "Analytical Thinking", "THK-01"),
]
for b, comp, code in BEHAVIORS:
    assert COMPETENCIES[code][0] == comp, (b, comp, code)
# Behaviors averaged into "Leadership + Agility Avg"
LEADERSHIP_AGILITY = {"Influence in group", "Conflict handling", "Guiding others", "Presence / authority",
                      "Decision quality", "Initiative", "Ownership", "Emotional regulation",
                      "Openness to feedback", "Learning speed", "Applies feedback"}


def potential_note(avg, la, growth_so_far, participation):
    if avg >= 3.9 and la >= 3.8:
        return "High potential – ready for stretch assignments"
    if avg >= 3.5:
        return "Strong potential – needs more exposure"
    if growth_so_far >= 0.5 and avg < 3.2:
        return "Emerging – improving fast"
    if avg < 2.5 and participation <= 2:
        return "Watch – low engagement"
    if avg < 2.5:
        return "Risk – struggles under pressure"
    return "Solid – steady contributor"


NOTE_TEMPLATES = [
    "Strongest on {s}; {w} still needs work.",
    "Reliable {s}. Tends to hold back on {w}.",
    "Good energy in group tasks. Next focus: {w}.",
    "Clear progress in {s}. Push for more {w}.",
    "Solid day overall; {w} was the weak spot.",
    "Stood out for {s}; coach on {w} before next session.",
]

obs_rows = []
for e, s, status, attendance in observed_enrolments:
    key = s["key"]
    title, ttype, mode, hours, days = COURSES[key]
    leniency = TRAINERS[s["trainer"]]
    day_dates = [s["start"]]
    while len(day_dates) < days:
        day_dates.append(add_workdays(day_dates[-1], 1))
    absent_day = rng.randrange(days) if attendance == "Partial" else None
    first_avg = None
    for d_idx, d in enumerate(day_dates):
        if d > TODAY:
            break
        if attendance == "No-Show" or d_idx == absent_day:
            obs_rows.append([s["code"], ttype, e["email"], e["name"], e["department"], e["role"], s["batch"], d_idx + 1, d,
                             title, s["trainer"], "Absent", None, None, None, None] + [None] * len(BEHAVIORS))
            continue
        scores = []
        for b, comp, code in BEHAVIORS:
            lvl = e["latent"][code]
            x = 2.6 + 0.4 * (lvl - 2.5) + 0.12 * e["grade"] + e["growth"] * d_idx + leniency + rng.gauss(0, 0.45)
            if e["archetype"] == "struggling" and b in ("Participation", "Initiative", "Influence in group"):
                x -= 0.6
            scores.append(clamp(rnd(x), 1, 5))
        avg = round(sum(scores) / len(scores), 2)
        la_scores = [sc for (b, _, _), sc in zip(BEHAVIORS, scores) if b in LEADERSHIP_AGILITY]
        la = round(sum(la_scores) / len(la_scores), 2)
        if first_avg is None:
            first_avg = avg
        ranked = sorted(zip(scores, [b for b, _, _ in BEHAVIORS], [rng.random() for _ in BEHAVIORS]))
        weak, strong = ranked[0][1].lower(), ranked[-1][1].lower()
        participation = scores[[b for b, _, _ in BEHAVIORS].index("Participation")]
        note = potential_note(avg, la, avg - first_avg, participation)
        assessor = rng.choice(NOTE_TEMPLATES).format(s=strong, w=weak)
        obs_rows.append([s["code"], ttype, e["email"], e["name"], e["department"], e["role"], s["batch"], d_idx + 1, d,
                         title, s["trainer"], "Present", avg, la, note, assessor] + scores)

# --------------------------------------------------------------------------------------
# Table 8: learning plans (only for people whose report has been delivered)
# --------------------------------------------------------------------------------------
PRIORITY_W = {"High": 3, "Medium": 2, "Low": 1}
plan_rows = []
for e in employees:
    if e["report"] != "Delivered":
        continue
    gaps = []
    for code in e["codes"]:
        req = required_level(code, e["grade"], e["tech_codes"])
        gap = req - e["latent"][code]
        if gap > 0:
            gaps.append((gap * PRIORITY_W[priority(code, e["grade"], e["tech_codes"])], rng.random(), code))
    gaps.sort(reverse=True)
    n_seg = rng.choice([3, 4, 4, 5, 5]) if e["grade"] < 4 else rng.choice([4, 5, 5])
    segments = []
    for _, _, code in gaps[:n_seg]:
        segments.append((rng.choice(COMPETENCIES[code][4]), rng.choice([3, 4, 4.5, 5, 6, 6, 8, 10, 12])))
    # people with few gaps get stretch work aimed at the next role
    stretch_codes = ["LEA-02", "COM-04", "THK-03", "LEA-01", "COM-03"]
    rng.shuffle(stretch_codes)
    while len(segments) < n_seg - 1 and stretch_codes:
        code = stretch_codes.pop()
        segments.append((f"Stretch: {COMPETENCIES[code][4][0]}", rng.choice([6, 8, 10])))
    if e["current_cert"] and e["current_cert"][1] not in ("Achieved",) and len(segments) < 5:
        segments.append((f"{e['current_cert'][0]} exam preparation", rng.choice([20, 24, 30, 40])))
    segments = segments[:5]
    need, need_h = DEPARTMENT_NEEDS.get(e["department"], ("Leading security culture across the division", 8))
    row = [e["name"], e["email"]]
    for i in range(5):
        if i < len(segments):
            row += [segments[i][0], fmt_hours(segments[i][1])]
        else:
            row += [None, None]
    row += [need, fmt_hours(need_h)]
    plan_rows.append(row)

# --------------------------------------------------------------------------------------
# Assemble the 8 sheets with the original column names
# --------------------------------------------------------------------------------------
T1_COLS = ["User Name", "User Email", "Job Title", "Job Title Extended", "Section", "Department", "Division",
           "Line Manager", "Self-Assessment", "Validation Interview", "Validation Interview Report",
           "Cross-Functional Training"]
t1 = [[e["name"], e["email"], short_title(e["role"]), e["extended"], e["section"], e["department"], e["division"],
       e["line_manager"][0], e["sa"], e["vi"], e["report"], e["cft"]] for e in employees]

T2_COLS = ["Team Name", "Team Manager Name", "Team Manager Email", "Line Manager Name", "Line Manager Email",
           "Person Name", "Person Email", "Job Role Family", "Job Role", "Framework", "Category", "Competency",
           "Competency Code", "Level", "Required Level", "Competency Gap", "Minimum Level", "Maximum Level", "Priority"]
t2 = []
for e in employees:
    if e["sa"] != "Completed":
        continue
    for code in e["codes"]:
        lo, hi = scale(code)
        level = e["latent"][code]
        if e["vi"] != "Completed" and rng.random() < 0.35:  # self-rated only: a little optimistic
            level = clamp(level + 1, lo, hi)
        req = required_level(code, e["grade"], e["tech_codes"])
        fw, cat = CATEGORIES[code[:3]]
        t2.append([e["team_name"], e["team_manager"][0], e["team_manager"][1], e["line_manager"][0], e["line_manager"][1],
                   e["name"], e["email"], e["family"], e["role"], fw, cat, COMPETENCIES[code][0], code,
                   level, req, req - level, lo, hi, priority(code, e["grade"], e["tech_codes"])])

T3_COLS = ["Framework Name", "Category Name", "Category Code", "Competency Minimum Level", "Competency Maximum Level",
           "Competency Code", "Competency Name", "Competency Description", "Supporting Text"]
t3 = []
for code, (name, desc, support, _, _) in COMPETENCIES.items():
    fw, cat = CATEGORIES[code[:3]]
    lo, hi = scale(code)
    t3.append([fw, cat, code[:3], lo, hi, code, name, desc, support])

T4_COLS = ["Framework Name", "Category Name", "Category Code", "Competency Code", "Competency Name", "Competency Behavior",
           "Competency Behavior Code", "CompetencyLevel", "Framework ExternalID", "Competency Category ExternalId",
           "Competency ExternalID", "Competency Behavior ExternalID"]
FW_EXT = {LEAD_FW: "FW-1001", TECH_FW: "FW-1002"}
CAT_EXT = {c: f"CAT-{2001 + i}" for i, c in enumerate(CATEGORIES)}
t4, beh_seq = [], 40001
for i, (code, (name, _, _, phrases, _)) in enumerate(COMPETENCIES.items()):
    fw, cat = CATEGORIES[code[:3]]
    for level in range(1, 6):
        for b_i, phrase in enumerate(phrases, start=1):
            t4.append([fw, cat, code[:3], code, name, LEVEL_TEMPLATES[level].format(p=phrase),
                       f"{code}-L{level}-B{b_i}", level, FW_EXT[fw], CAT_EXT[code[:3]], f"CMP-{3001 + i}",
                       f"BEH-{beh_seq}"])
            beh_seq += 1

T5_COLS = ["Training Record Code", "Training Type", "Employee Email", "Training Start Date", "Training End Date",
           "Training Title", "Attendance", "Status", "Mode", "Hours"]
training_rows.sort(key=lambda r: (r["session"]["start"], r["session"]["code"], r["email"]))
t5 = [[r["session"]["code"], r["ttype"], r["email"], r["session"]["start"], r["session"]["end"], r["title"],
       r["attendance"] or None, r["status"], r["mode"], r["hours"]] for r in training_rows]

T6_COLS = ["Training Record Code", "Training Type", "Employee Email", "Employee Name", "Department", "Role", "Batch",
           "Training Day", "Training Date", "Training Title", "Trainer Name", "Attendance", "Overall Day Avg",
           "Leadership + Agility Avg", "Potential / Risk Note", "Assessor Notes"] + [f"{b} | {c}" for b, c, _ in BEHAVIORS]
obs_rows.sort(key=lambda r: (r[8], r[0], r[3]))
t6 = obs_rows

T7_COLS = ["IDP Cycle", "User Email", "Prof Cert", "Prof Cert Status", "Prof Cert Description", "AI Cert",
           "AI Cert Status", "AI Cert Description"]
t7 = idp_rows

T8_COLS = ["Name", "Email"]
for i in range(1, 6):
    T8_COLS += [f"Segment {i}", f"Estimated Time {i}"]
T8_COLS += ["Part 2: Department Training Needs", "Estimated Time Part 2"]
t8 = plan_rows

SHEETS = [
    ("Employees", T1_COLS, t1),
    ("Competency Gaps", T2_COLS, t2),
    ("Competencies", T3_COLS, t3),
    ("Competency Behaviors", T4_COLS, t4),
    ("Training Records", T5_COLS, t5),
    ("Training Day Scores", T6_COLS, t6),
    ("IDP Certifications", T7_COLS, t7),
    ("Learning Plans", T8_COLS, t8),
]

wb = Workbook()
wb.remove(wb.active)
header_font = Font(bold=True, color="FFFFFF")
header_fill = PatternFill("solid", fgColor="1F3A5F")
for title, cols, rows in SHEETS:
    for r in rows:
        assert len(r) == len(cols), (title, len(r), len(cols))
    ws = wb.create_sheet(title)
    ws.append(cols)
    for r in rows:
        ws.append(r)
    for c in ws[1]:
        c.font, c.fill = header_font, header_fill
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, col in enumerate(cols, start=1):
        sample = [len(str(col))] + [len(str(r[i - 1])) for r in rows[:400] if r[i - 1] is not None]
        ws.column_dimensions[get_column_letter(i)].width = min(60, max(10, max(sample) + 2))
        if any(isinstance(r[i - 1], dt.date) for r in rows[:50]):
            for row in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                row[0].number_format = "yyyy-mm-dd"

OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print(f"Wrote {OUT}")
for title, cols, rows in SHEETS:
    print(f"  {title:<22} {len(rows):>6} rows  {len(cols):>3} columns")
